from src.wis.ccs_website_data import fetch_all_ccs_frameworks
from azure.storage.blob import ContainerClient, ExponentialRetry, ContentSettings
import os
from dotenv import load_dotenv
import io
from docx import Document


load_dotenv()

ccs_frameworks = fetch_all_ccs_frameworks()
container_client = ContainerClient.from_connection_string(
    conn_str=os.getenv("BLOB_CONNECTION_STRING"),
    container_name=os.getenv("BLOB_CONTAINER_NAME"),
    retry_policy=ExponentialRetry(initial_backoff=2, retry_total=5),
)

for index, row in ccs_frameworks.iterrows():
    # create a Word Document in memory
    doc = Document()

    # add a styled Title
    doc.add_heading(str(row["title"]), 0)

    # add structured sections
    doc.add_heading('Framework Details', level=1)
    p = doc.add_paragraph()
    p.add_run('RM Number: ').bold = True
    p.add_run(str(row["rm_number"]))

    doc.add_heading('Description', level=2)
    doc.add_paragraph(str(row["description"]))

    doc.add_heading('Summary', level=2)
    doc.add_paragraph(str(row["summary"]))

    doc.add_heading('benefits', level=2)
    doc.add_paragraph(str(row["benefits"]))

    doc.add_heading('how_to_buy', level=2)
    doc.add_paragraph(str(row["how_to_buy"]))

    # save to a Bytes buffer instead of a file
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)  # reset buffer pointer to the beginning

    # upload to Azure
    azure_file_name = f"{row['rm_number']}_page_content.docx"
    blob_client = container_client.get_blob_client(azure_file_name)

    print(f"Uploading {azure_file_name}...")

    blob_client.upload_blob(
        data=doc_io,
        overwrite=True,
        metadata={"rm_number": str(row["rm_number"])},
        content_settings=ContentSettings(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            content_disposition=f'attachment; filename={azure_file_name}'
        )
    )
